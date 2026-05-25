from fastapi import FastAPI, HTTPException, Request, status, UploadFile, File
from fastapi.responses import JSONResponse
from groq import Groq
import os
import json
import traceback
import io
import PyPDF2
from pptx import Presentation
from dotenv import load_dotenv

load_dotenv()
app = FastAPI()
client = Groq()

@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception):
    traceback.print_exc()
    return JSONResponse(
        status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
        content={"detail": "Internal Server Error", "error": str(exc)},
    )

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    text = ""
    try:
        if filename.lower().endswith('.pdf'):
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
                    
        elif filename.lower().endswith('.pptx'):
            prs = Presentation(io.BytesIO(file_bytes))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        else:
            raise ValueError("Unsupported file type. Please upload a PDF or PPTX.")
    except Exception as e:
        print(f"Error extracting text: {e}")
        raise ValueError("Failed to read file content.")
        
    return text.strip()

@app.post("/generate")
async def generate_materials(file: UploadFile = File(...)):
    # 1. Read and extract text from the uploaded file
    file_bytes = await file.read()
    try:
        text_content = extract_text_from_file(file_bytes, file.filename)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    if not text_content:
        raise HTTPException(status_code=400, detail="Could not extract any text from the file.")

    # 2. Send the extracted text to Llama
    system_prompt = """
    You are an expert educational assistant. Analyze the provided text and generate study materials.
    You must return a strictly valid JSON object with two keys:
    1. 'flashcards': an array of objects, each with 'front' (the question/concept) and 'back' (the answer).
    2. 'quiz': an array of objects, each with 'question', 'options' (array of 4 strings), and 'correct_answer' (the exact string of the correct option).
    """

    chat_completion = client.chat.completions.create(
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": text_content[:15000]} # Limit characters to avoid token overload
        ],
        model="llama-3.3-70b-versatile"
        response_format={"type": "json_object"},
        temperature=0.3
    )

    raw_content = chat_completion.choices[0].message.content
    cleaned_content = raw_content.replace("```json", "").replace("```", "").strip()
    
    try:
        result = json.loads(cleaned_content)
        return result
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="AI returned invalid JSON format.")